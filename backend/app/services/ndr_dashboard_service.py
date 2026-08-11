from __future__ import annotations

from collections import Counter, defaultdict
from datetime import date, datetime, timedelta
from io import BytesIO
import math
import re
from typing import Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session
from starlette.responses import StreamingResponse

from app.models.entities import NdrTrackingRecord


REQUIRED_COLUMNS = ["OrderNo", "PincodeZone", "Shipment", "OUR EDD", "Current status"]
AGE_BUCKETS = [
    ("01-05", 1, 5),
    ("06-10", 6, 10),
    ("11-15", 11, 15),
    ("16-20", 16, 20),
    ("21-25", 21, 25),
    ("26-31", 26, 31),
]


def _raw_value(raw: dict, column: str):
    value = raw.get(column)
    if value is None:
        return None
    cleaned = str(value).strip()
    return cleaned or None


def _raw_first(raw: dict, columns: list[str]):
    for column in columns:
        value = _raw_value(raw, column)
        if value:
            return value
    normalized = {re.sub(r"[^a-z0-9]", "", key.lower()): value for key, value in raw.items()}
    for column in columns:
        value = normalized.get(re.sub(r"[^a-z0-9]", "", column.lower()))
        if value is not None and str(value).strip():
            return str(value).strip()
    return None


def _normalize(value: str | None) -> str:
    return re.sub(r"\s+", " ", str(value or "").strip()).lower()


def _status_group(status: str | None) -> str:
    normalized = _normalize(status)
    if "deliver" in normalized:
        return "Delivered"
    if any(token in normalized for token in ["ship", "transit", "ofd", "out for delivery", "pickup", "manifest"]):
        return "Shipped / In Transit"
    if any(token in normalized for token in ["packed", "pack"]):
        return "Packed"
    if any(token in normalized for token in ["process", "pending", "ndr", "hold", "attempt", "rto"]):
        return "Pending"
    return "Other Status" if normalized else "Other Status"


def _courier_group(courier: str | None) -> str:
    cleaned = str(courier or "").strip() or "Unknown"
    normalized = _normalize(cleaned)
    if "dtdc" in normalized and ("5g" in normalized or "5b" in normalized):
        return "DTDC"
    return cleaned


def _is_delivered(status: str | None) -> bool:
    return _status_group(status) == "Delivered"


def _is_shipped(status: str | None) -> bool:
    return _status_group(status) == "Shipped / In Transit"


def _is_not_shipped(status: str | None) -> bool:
    return not _is_delivered(status) and not _is_shipped(status)


def _parse_date(value) -> date | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    if not text or text.lower() in {"nan", "nat", "none"}:
        return None
    for dayfirst in (False, True):
        try:
            parsed = datetime.strptime(text[:10], "%Y-%m-%d").date()
            return parsed
        except ValueError:
            pass
        for fmt in ("%d-%m-%Y", "%d/%m/%Y", "%d.%m.%Y", "%m/%d/%Y"):
            try:
                return datetime.strptime(text[:10], fmt).date()
            except ValueError:
                continue
        if not dayfirst:
            continue
    try:
        import pandas as pd

        parsed = pd.to_datetime(text, errors="coerce", dayfirst=True)
        if not pd.isna(parsed):
            return parsed.date()
    except Exception:
        return None
    return None


def _pct(value: int, total: int) -> float:
    return round((value / total) * 100, 1) if total else 0.0


def _date_window(name: str | None) -> tuple[date | None, date | None]:
    today = date.today()
    if name == "today":
        return today, today
    if name == "yesterday":
        day = today - timedelta(days=1)
        return day, day
    if name == "last7":
        return today - timedelta(days=6), today
    if name == "last15":
        return today - timedelta(days=14), today
    if name == "last30":
        return today - timedelta(days=29), today
    return None, None


def _base_orders(db: Session) -> tuple[list[dict], list[str]]:
    rows = (
        db.query(NdrTrackingRecord.id, NdrTrackingRecord.upload_filename, NdrTrackingRecord.row_number, NdrTrackingRecord.raw_data)
        .order_by(NdrTrackingRecord.id.asc())
        .all()
    )
    if not rows:
        return [], REQUIRED_COLUMNS.copy()

    available = set()
    source_rows: dict[tuple[str, int, str], dict] = {}
    for row_id, upload_filename, row_number, raw in rows:
        raw = raw or {}
        available.update(raw.keys())
        order_no = _raw_value(raw, "OrderNo")
        if not order_no:
            continue
        item = {
            "id": row_id,
            "order_no": order_no,
            "zone": _raw_value(raw, "PincodeZone") or "Unknown",
            "courier": _courier_group(_raw_value(raw, "Shipment")),
            "warehouse": _raw_first(raw, ["Warehouse", "Store"]) or "Unknown",
            "shipping_date": _parse_date(_raw_first(raw, ["ShippingDate", "Shipping Date"])),
            "cx_name": _raw_first(raw, ["Cx Name", "CX Name", "Customer Name", "Customer"]),
            "mobile_no": _raw_first(raw, ["Mobile No", "Mobile", "Phone"]),
            "docket_no": _raw_first(raw, ["Docketno", "Docket No", "Docket", "AWB"]),
            "edd": _parse_date(_raw_value(raw, "OUR EDD")),
            "status": _raw_value(raw, "Current status") or "Unknown",
        }
        source_rows[(upload_filename or "", row_number or 0, order_no)] = item

    missing = [column for column in REQUIRED_COLUMNS if column not in available]
    return list(source_rows.values()), missing


def _apply_filters(orders: list[dict], params: dict) -> list[dict]:
    date_filter = params.get("date_filter")
    start, end = _date_window(date_filter)
    if params.get("start_date"):
        start = _parse_date(params.get("start_date"))
    if params.get("end_date"):
        end = _parse_date(params.get("end_date"))

    courier = params.get("courier")
    zone = params.get("zone")
    status = params.get("status")
    edd_status = params.get("edd_status")
    search = _normalize(params.get("search"))
    today = date.today()

    filtered = []
    for order in orders:
        if courier and courier != "all" and order["courier"] != courier:
            continue
        if zone and zone != "all" and order["zone"] != zone:
            continue
        if status and status != "all" and order["status"] != status:
            continue
        if search and search not in _normalize(order["order_no"]):
            continue
        filter_date = order.get("shipping_date") or order["edd"]
        if start and (not filter_date or filter_date < start):
            continue
        if end and (not filter_date or filter_date > end):
            continue
        current_edd = _edd_status(order, today)
        if edd_status and edd_status != "all" and current_edd != edd_status:
            continue
        filtered.append(order)
    return filtered


def _edd_status(order: dict, today: date | None = None) -> str:
    today = today or date.today()
    if _is_delivered(order["status"]):
        return "Delivered"
    if order["edd"] and order["edd"] < today:
        return "EDD Expired"
    return "EDD Remaining"


def _pending_days(order: dict, today: date | None = None) -> int:
    today = today or date.today()
    if _is_delivered(order["status"]) or not order["edd"]:
        return 0
    return max((today - order["edd"]).days, 0)


def _count_group(orders: Iterable[dict]) -> dict:
    orders = list(orders)
    today = date.today()
    total = len(orders)
    delivered = sum(1 for order in orders if _is_delivered(order["status"]))
    shipped = sum(1 for order in orders if _is_shipped(order["status"]))
    not_shipped = sum(1 for order in orders if _is_not_shipped(order["status"]))
    pending = total - delivered
    edd_expired = sum(1 for order in orders if _edd_status(order, today) == "EDD Expired")
    edd_remaining = sum(1 for order in orders if _edd_status(order, today) == "EDD Remaining")
    return {
        "total": total,
        "delivered": delivered,
        "shipped": shipped,
        "not_shipped": not_shipped,
        "pending": pending,
        "edd_expired": edd_expired,
        "edd_remaining": edd_remaining,
        "delivery_percentage": _pct(delivered, total),
    }


def _group_table(orders: list[dict], key: str, label: str) -> list[dict]:
    grouped: dict[str, list[dict]] = defaultdict(list)
    for order in orders:
        grouped[order[key] or "Unknown"].append(order)
    table = []
    for name, group_orders in grouped.items():
        metrics = _count_group(group_orders)
        table.append({label: name, **metrics})
    return sorted(table, key=lambda row: row["pending"], reverse=True)


def _status_overview(orders: list[dict]) -> list[dict]:
    total = len(orders)
    counts = Counter(_status_group(order["status"]) for order in orders)
    return [{"status": name, "count": count, "percentage": _pct(count, total)} for name, count in counts.most_common()]


def _shipping_bucket_label(shipping_date: date | None) -> str:
    if not shipping_date:
        return "No Shipping Date"
    for label, start, end in AGE_BUCKETS:
        if start <= shipping_date.day <= end:
            return f"{label} {shipping_date.strftime('%b %Y')}"
    return f"Other {shipping_date.strftime('%b %Y')}"


def _ageing(orders: list[dict]) -> list[dict]:
    grouped: dict[str, list[dict]] = defaultdict(list)
    for order in orders:
        grouped[_shipping_bucket_label(order.get("shipping_date"))].append(order)
    result = []
    for label, group_orders in grouped.items():
        bucket_orders = list(group_orders)
        result.append({"bucket": label, **_count_group(bucket_orders)})
    return sorted(result, key=lambda row: row["bucket"])


def _pending_order_analysis_detail(orders: list[dict]) -> list[dict]:
    grouped: dict[tuple[str, str, str, str], list[dict]] = defaultdict(list)
    for order in orders:
        grouped[
            (
                _shipping_bucket_label(order.get("shipping_date")),
                order["courier"] or "Unknown",
                order.get("warehouse") or "Unknown",
                order["zone"] or "Unknown",
            )
        ].append(order)
    rows = []
    for (bucket, courier, warehouse, zone), group_orders in grouped.items():
        metrics = _count_group(group_orders)
        rows.append(
            {
                "shipping_date_range": bucket,
                "courier": courier,
                "warehouse": warehouse,
                "zone": zone,
                **metrics,
            }
        )
    return sorted(rows, key=lambda row: (row["shipping_date_range"], row["courier"], row["warehouse"], row["zone"]))


def _critical_order_details(orders: list[dict]) -> list[dict]:
    today = date.today()
    rows = []
    for order in orders:
        if _is_delivered(order["status"]) or _edd_status(order, today) != "EDD Expired":
            continue
        rows.append(
            {
                "order_no": order["order_no"],
                "cx_name": order.get("cx_name"),
                "shipment": order["courier"],
                "phone_no": order.get("mobile_no"),
                "docket_no": order.get("docket_no"),
                "warehouse": order.get("warehouse"),
                "zone": order["zone"],
                "shipping_date": order["shipping_date"].isoformat() if order.get("shipping_date") else None,
                "our_edd": order["edd"].isoformat() if order["edd"] else None,
                "current_status": order["status"],
                "edd_status": _edd_status(order, today),
                "pending_days": _pending_days(order, today),
            }
        )
    return sorted(rows, key=lambda row: (-row["pending_days"], row["order_no"]))


def _pending_orders(orders: list[dict]) -> list[dict]:
    today = date.today()
    rows = []
    for order in orders:
        if _is_delivered(order["status"]):
            continue
        rows.append(
            {
                "order_no": order["order_no"],
                "courier": order["courier"],
                "zone": order["zone"],
                "current_status": order["status"],
                "our_edd": order["edd"].isoformat() if order["edd"] else None,
                "shipping_date": order["shipping_date"].isoformat() if order.get("shipping_date") else None,
                "edd_status": _edd_status(order, today),
                "pending_days": _pending_days(order, today),
            }
        )
    return sorted(rows, key=lambda row: (row["edd_status"] != "EDD Expired", -row["pending_days"], row["order_no"]))


def ndr_dashboard(db: Session, params: dict | None = None) -> dict:
    params = params or {}
    orders, missing = _base_orders(db)
    if missing:
        return {"valid": False, "missing_columns": missing, "message": f"Required column missing: {missing[0]}. Please upload an Excel file containing the required column."}

    filtered = _apply_filters(orders, params)
    metrics = _count_group(filtered)
    status = _status_overview(filtered)
    courier = _group_table(filtered, "courier", "courier")
    zone = _group_table(filtered, "zone", "zone")
    ageing = _ageing(filtered)
    pending_analysis_detail = _pending_order_analysis_detail(filtered)
    pending = _pending_orders(filtered)
    critical_orders = _critical_order_details(filtered)
    edd = [
        {"name": "EDD Delivered", "value": metrics["delivered"], "percentage": _pct(metrics["delivered"], metrics["total"])},
        {"name": "EDD Remaining", "value": metrics["edd_remaining"], "percentage": _pct(metrics["edd_remaining"], metrics["total"])},
        {"name": "EDD Expired", "value": metrics["edd_expired"], "percentage": _pct(metrics["edd_expired"], metrics["total"])},
    ]
    top_courier = courier[0] if courier else None
    top_zone = zone[0] if zone else None
    critical_30 = max(ageing, key=lambda row: row["pending"], default={"pending": 0, "bucket": "No Shipping Date"})
    insights = [
        f"Total {metrics['total']:,} NDR rows from the uploaded file are being reviewed, with {metrics['delivered']:,} delivered and {metrics['pending']:,} requiring action.",
        f"{metrics['edd_expired']:,} pending orders have crossed OUR EDD and need immediate follow-up.",
    ]
    if top_courier:
        insights.append(f"{top_courier['courier']} has the highest pending load with {top_courier['pending']:,} orders.")
    if top_zone:
        insights.append(f"{top_zone['zone']} is the highest pending zone with {top_zone['pending']:,} orders.")
    if critical_30["pending"]:
        insights.append(f"{critical_30['bucket']} has {critical_30['pending']:,} pending orders and should be reviewed.")

    alerts = []
    if metrics["edd_expired"]:
        alerts.append({"level": "critical", "message": f"{metrics['edd_expired']:,} orders have crossed OUR EDD."})
    not_shipped_count = metrics.get("not_shipped", 0)
    if not_shipped_count:
        alerts.append({"level": "attention", "message": f"{not_shipped_count:,} orders are not shipped yet as per Current status."})
    if top_zone:
        alerts.append({"level": "monitor", "message": f"{top_zone['zone']} has the highest pending volume."})
    if top_courier:
        alerts.append({"level": "monitor", "message": f"{top_courier['courier']} needs close review for pending orders."})

    return {
        "valid": True,
        "report_date": date.today().isoformat(),
        "filters": {
            "couriers": sorted({order["courier"] for order in orders}),
            "zones": sorted({order["zone"] for order in orders}),
            "statuses": sorted({order["status"] for order in orders}),
        },
        "kpis": {
            "total_orders": metrics["total"],
            "delivered": metrics["delivered"],
            "delivered_percentage": _pct(metrics["delivered"], metrics["total"]),
            "shipped": metrics["shipped"],
            "shipped_percentage": _pct(metrics["shipped"], metrics["total"]),
            "pending": metrics["pending"],
            "pending_percentage": _pct(metrics["pending"], metrics["total"]),
            "edd_expired": metrics["edd_expired"],
            "edd_expired_percentage": _pct(metrics["edd_expired"], metrics["pending"]),
            "edd_remaining": metrics["edd_remaining"],
            "edd_remaining_percentage": _pct(metrics["edd_remaining"], metrics["pending"]),
            "delivery_percentage": metrics["delivery_percentage"],
        },
        "status_overview": status,
        "courier_performance": courier,
        "zone_performance": zone,
        "pending_ageing": ageing,
        "pending_order_analysis_detail": pending_analysis_detail,
        "edd_performance": edd,
        "pending_orders": pending,
        "critical_orders": critical_orders,
        "insights": insights,
        "alerts": alerts,
        "search_result": _search_order(filtered, params.get("order_search")),
    }


def _search_order(orders: list[dict], order_no: str | None) -> dict | None:
    if not order_no:
        return None
    lookup = _normalize(order_no)
    today = date.today()
    for order in orders:
        if _normalize(order["order_no"]) == lookup:
            return {
                "order_no": order["order_no"],
                "courier": order["courier"],
                "zone": order["zone"],
                "current_status": order["status"],
                "our_edd": order["edd"].isoformat() if order["edd"] else None,
                "order_age": _pending_days(order, today),
                "edd_status": _edd_status(order, today),
            }
    return None


def _append_table(ws, title: str, headers: list[str], rows: list[list], start_row: int) -> int:
    ws.cell(start_row, 1, title)
    ws.cell(start_row, 1).font = Font(bold=True, size=14, color="FFFFFF")
    ws.cell(start_row, 1).fill = PatternFill("solid", fgColor="1F4E78")
    start_row += 1
    for col, header in enumerate(headers, 1):
        cell = ws.cell(start_row, col, header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4472C4")
    for row_index, row in enumerate(rows, start_row + 1):
        for col, value in enumerate(row, 1):
            ws.cell(row_index, col, value)
    end_row = start_row + len(rows)
    if rows:
        ws.auto_filter.ref = f"A{start_row}:{get_column_letter(len(headers))}{end_row}"
    for column in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(column)].width = 18
    thin = Side(style="thin", color="D9E2F3")
    for row in ws.iter_rows(min_row=start_row, max_row=end_row, max_col=len(headers)):
        for cell in row:
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
            cell.alignment = Alignment(vertical="top")
    return end_row + 3


def excel_report(db: Session, params: dict | None = None) -> StreamingResponse:
    data = ndr_dashboard(db, params)
    wb = Workbook()
    ws = wb.active
    ws.title = "Executive Summary"
    kpis = data.get("kpis", {})
    _append_table(
        ws,
        "NDR Management Report",
        ["Metric", "Value"],
        [
            ["Report Date", data.get("report_date")],
            ["Total Orders", kpis.get("total_orders", 0)],
            ["Delivered", kpis.get("delivered", 0)],
            ["Shipped", kpis.get("shipped", 0)],
            ["Pending", kpis.get("pending", 0)],
            ["EDD Expired", kpis.get("edd_expired", 0)],
            ["EDD Remaining", kpis.get("edd_remaining", 0)],
            ["Delivery %", kpis.get("delivery_percentage", 0)],
        ],
        1,
    )
    sheets = [
        ("Order Status", ["Status", "Order Count", "Percentage"], [[r["status"], r["count"], r["percentage"]] for r in data.get("status_overview", [])]),
        ("Courier Performance", ["Courier", "Total", "Delivered", "Shipped", "Pending", "EDD Expired", "EDD Remaining", "Delivery %"], [[r["courier"], r["total"], r["delivered"], r["shipped"], r["pending"], r["edd_expired"], r["edd_remaining"], r["delivery_percentage"]] for r in data.get("courier_performance", [])]),
        ("Zone Performance", ["Zone", "Total", "Delivered", "Shipped", "Pending", "EDD Expired", "EDD Remaining"], [[r["zone"], r["total"], r["delivered"], r["shipped"], r["pending"], r["edd_expired"], r["edd_remaining"]] for r in data.get("zone_performance", [])]),
        ("Pending Order Analysis", ["Shipping Date Range", "Courier", "Store / Warehouse", "Zone", "Total Orders", "Shipped", "Not Shipped", "Delivered", "Pending", "EDD Expired", "EDD Remaining"], [[r["shipping_date_range"], r["courier"], r["warehouse"], r["zone"], r["total"], r["shipped"], r["not_shipped"], r["delivered"], r["pending"], r["edd_expired"], r["edd_remaining"]] for r in data.get("pending_order_analysis_detail", [])]),
        ("Critical Attention", ["Order No", "Cx Name", "Shipment", "Phone No", "Docket No", "Store / Warehouse", "Zone", "Shipping Date", "OUR EDD", "Current Status", "EDD Status", "Pending Days"], [[r["order_no"], r["cx_name"], r["shipment"], r["phone_no"], r["docket_no"], r["warehouse"], r["zone"], r["shipping_date"], r["our_edd"], r["current_status"], r["edd_status"], r["pending_days"]] for r in data.get("critical_orders", [])]),
        ("Critical Summary", ["Level", "Alert"], [[r["level"].title(), r["message"]] for r in data.get("alerts", [])]),
        ("Management Insights", ["Insight"], [[insight] for insight in data.get("insights", [])]),
        ("Pending Orders", ["Order No", "Courier", "Zone", "Current Status", "Shipping Date", "OUR EDD", "EDD Status", "Pending Days"], [[r["order_no"], r["courier"], r["zone"], r["current_status"], r["shipping_date"], r["our_edd"], r["edd_status"], r["pending_days"]] for r in data.get("pending_orders", [])]),
    ]
    for title, headers, rows in sheets:
        sheet = wb.create_sheet(title)
        sheet.freeze_panes = "A3"
        _append_table(sheet, title, headers, rows, 1)

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="ndr-management-report.xlsx"'},
    )


def _pdf_escape(text: str) -> str:
    return str(text).replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def pdf_report(db: Session, params: dict | None = None) -> StreamingResponse:
    data = ndr_dashboard(db, params)
    kpis = data.get("kpis", {})
    lines = [
        "NDR MANAGEMENT REPORT",
        f"Reporting Date: {data.get('report_date')}",
        "",
        f"Total Orders: {kpis.get('total_orders', 0):,}",
        f"Delivered: {kpis.get('delivered', 0):,} ({kpis.get('delivered_percentage', 0)}%)",
        f"Shipped / In Transit: {kpis.get('shipped', 0):,} ({kpis.get('shipped_percentage', 0)}%)",
        f"Pending Orders: {kpis.get('pending', 0):,} ({kpis.get('pending_percentage', 0)}%)",
        f"EDD Expired: {kpis.get('edd_expired', 0):,}",
        f"EDD Remaining: {kpis.get('edd_remaining', 0):,}",
        "",
        "Management Insights",
        *data.get("insights", []),
        "",
        "Critical Attention Required",
        *[alert["message"] for alert in data.get("alerts", [])],
    ]
    content_lines = ["BT", "/F1 12 Tf", "50 790 Td"]
    for index, line in enumerate(lines[:45]):
        if index:
            content_lines.append("0 -18 Td")
        content_lines.append(f"({_pdf_escape(line)}) Tj")
    content_lines.append("ET")
    stream = "\n".join(content_lines).encode("latin-1", "replace")
    objects = [
        b"1 0 obj << /Type /Catalog /Pages 2 0 R >> endobj",
        b"2 0 obj << /Type /Pages /Kids [3 0 R] /Count 1 >> endobj",
        b"3 0 obj << /Type /Page /Parent 2 0 R /MediaBox [0 0 612 842] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >> endobj",
        b"4 0 obj << /Type /Font /Subtype /Type1 /BaseFont /Helvetica >> endobj",
        f"5 0 obj << /Length {len(stream)} >> stream\n".encode() + stream + b"\nendstream endobj",
    ]
    pdf = BytesIO()
    pdf.write(b"%PDF-1.4\n")
    offsets = []
    for obj in objects:
        offsets.append(pdf.tell())
        pdf.write(obj + b"\n")
    xref = pdf.tell()
    pdf.write(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    for offset in offsets:
        pdf.write(f"{offset:010d} 00000 n \n".encode())
    pdf.write(f"trailer << /Root 1 0 R /Size {len(objects) + 1} >>\nstartxref\n{xref}\n%%EOF".encode())
    pdf.seek(0)
    return StreamingResponse(pdf, media_type="application/pdf", headers={"Content-Disposition": 'attachment; filename="ndr-management-report.pdf"'})
